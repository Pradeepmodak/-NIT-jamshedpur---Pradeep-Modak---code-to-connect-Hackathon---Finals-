"""
Generate exactly 7 slides for SentinelRisk using python-pptx.
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Add project root to path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from engine.config import OUTPUT_DIR, TICKERS

def create_deck():
    prs = Presentation()
    # 16:9
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # Dark theme colors
    bg_color = RGBColor(30, 30, 30)
    title_color = RGBColor(255, 255, 255)
    text_color = RGBColor(200, 200, 200)

    # Helper for blank slides
    blank_slide_layout = prs.slide_layouts[6]

    def add_slide(title_text, content_text="", notes_text=""):
        slide = prs.slides.add_slide(blank_slide_layout)
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1.5))
        tf = title_box.text_frame
        p = tf.add_paragraph()
        p.text = title_text
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = title_color
        
        # Content
        if content_text:
            content_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(5))
            tf2 = content_box.text_frame
            tf2.word_wrap = True
            for line in content_text.split('\n'):
                p = tf2.add_paragraph()
                p.text = line
                p.font.size = Pt(24)
                p.font.color.rgb = text_color
                
        # Notes
        if notes_text:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = notes_text
            
        return slide

    # Read backtest results
    bt_path = OUTPUT_DIR / "backtest.json"
    bt_data = {}
    if bt_path.exists():
        with open(bt_path, "r") as f:
            bt_data = json.load(f)
            
    dyn = bt_data.get("dynamic", {})
    ew = bt_data.get("equal_weight", {})

    # Slide 1: Problem and pitch
    add_slide(
        "SentinelRisk",
        "An AI/NLP financial risk engine and tactical index rebalancer.\n\n"
        "• The Problem: Markets react instantly to unstructured text (news, social media).\n"
        "• The Solution: SentinelRisk quantifies textual impact in real-time to tactically tilt index weights.\n"
        "• The Edge: Avoids major drawdowns from negative sentiment while capturing upside.",
        "SentinelRisk is a real-time AI/NLP financial risk engine. It ingests unstructured text, scores the sentiment and impact, and tactically rebalances an index to optimize risk-adjusted returns."
    )
    
    # Slide 2: Architecture
    add_slide(
        "Architecture",
        "1. Sources: NewsAPI & Kaggle Stock Tweets (Real-time / Replay)\n"
        "2. Engine: Ingestion → Common Schema\n"
        "3. NLP Pipeline: Entity Extraction (15 tickers) → FinBERT → Event Classifier\n"
        "4. Scoring: Impact Scorer (1-10) → API (FastAPI) & JSONL\n"
        "5. Module A: Tactical Index Rebalancer (EWMA → exp tilt → water-fill)\n"
        "6. Dashboard: Streamlit + Plotly",
        "Our architecture is modular. We ingest from News and Twitter. The NLP pipeline extracts entities, scores sentiment with FinBERT, and classifies events. The Impact Scorer aggregates this into a 1-10 signal, driving Module A's tactical rebalancing."
    )
    
    # Slide 3: NLP approach
    add_slide(
        "NLP & Scoring Approach",
        "• Entity Extractor: Alias dictionary + word-boundary regex.\n"
        "• Sentiment: ProsusAI/finbert (Domain-tuned, batched).\n"
        "• Event Classifier: Hybrid keyword-rules → DeBERTa zero-shot fallback.\n"
        "• Impact Formula:\n"
        "  Clip(Round(10 × (0.4×|sentiment| + 0.35×event_weight + 0.25×severity_kw))))",
        "We use FinBERT for sentiment and a hybrid approach for events, falling back to a zero-shot DeBERTa model if needed. The impact score combines sentiment magnitude, event type importance, and severity keywords."
    )

    # Slide 4: Module A
    add_slide(
        "Module A: Tactical Index Rebalancing",
        "• Tracks an EWMA of sentiment (configurable half-life) for 15 stocks.\n"
        "• Raw weight ∝ Base Weight × exp(k × EWMA_Sentiment).\n"
        "• Bounds: Iterative water-filling algorithm to enforce [2%, 15%] limits.\n"
        "• No Look-Ahead: Weights derived from T-1 signals applied to T returns.\n"
        "• Friction: Accounts for 5 bps transaction costs per unit turnover.",
        "Module A dynamically tilts a base equal-weight portfolio. We track sentiment using an EWMA, apply an exponential tilt, and enforce strict 2 to 15 percent bounds using iterative water-filling. Transaction costs are fully accounted for."
    )
    
    # Slide 5: Challenges and fixes
    add_slide(
        "Challenges and Solutions",
        "• Latency: Implemented batched sentiment inference.\n"
        "• Ambiguity: Multi-ticker matching and 'MARKET' fallback for general news.\n"
        "• Noise: Hybrid event classifier ignores irrelevant noise.\n"
        "• Reliability: Graceful fallbacks (Offline mode, keyword-sentiment, cached prices).\n"
        "• Look-Ahead Bias: Strict chronological separation in the backtest.",
        "We overcame several challenges. Latency was solved with batching. Entity ambiguity is handled gracefully. Offline reliability is a core feature, falling back to cached or synthetic data. The backtest strictly avoids look-ahead bias."
    )

    # Slide 6: Results
    ret_dyn = dyn.get('total_return_pct', 0)
    ret_ew = ew.get('total_return_pct', 0)
    vol_dyn = dyn.get('annualized_volatility_pct', 0)
    sharpe_dyn = dyn.get('sharpe_ratio', 0)
    dd_dyn = dyn.get('max_drawdown_pct', 0)
    
    add_slide(
        "Backtest Results (Replay Window)",
        f"Dynamic vs. Equal-Weight:\n\n"
        f"• Total Return: {ret_dyn:.2f}% (Dynamic) vs {ret_ew:.2f}% (EW)\n"
        f"• Ann. Volatility: {vol_dyn:.2f}%\n"
        f"• Sharpe Ratio: {sharpe_dyn:.4f}\n"
        f"• Max Drawdown: {dd_dyn:.2f}%\n"
        f"• Avg Turnover: {bt_data.get('avg_turnover_pct', 0):.2f}%",
        "Here are the backtest results. We see the dynamic index's performance compared to the equal-weight benchmark, showcasing the total return, volatility, Sharpe ratio, and maximum drawdown."
    )
    
    # Slide 7: Future work
    add_slide(
        "Future Work",
        "• Module B: Stress testing using event type triggers (Impact > 7).\n"
        "• Data Expansion: GDELT ingestion for macro geopolitical coverage.\n"
        "• Infrastructure: Streaming Kafka pipeline for real-time scale.\n"
        "• Models: Fine-tuned FinBERT on proprietary data.\n"
        "• Portfolio: Advanced convex optimization constraints.",
        "Looking ahead, we plan to implement Module B for stress testing, integrate GDELT for macro news, move to a Kafka streaming architecture, fine-tune our models, and add advanced portfolio optimization constraints."
    )

    out_path = Path("SentinelRisk.pptx")
    prs.save(str(out_path))
    print(f"Created {out_path}")

if __name__ == "__main__":
    create_deck()
